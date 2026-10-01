"""Extração incremental de documentos para ferramentas do agente.

Cada chamada devolve um trecho pequeno com posição de continuação. O arquivo
permanece em disco; não guardamos seu texto integral no histórico do modelo.
"""

import csv
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

TAMANHO_PADRAO = 2500
TAMANHO_MAXIMO = 50000
LIMITE_BUSCA_CARACTERES = 5_000_000
LIMITE_RESUMO = 10_000

EXTENSOES_TEXTO = {
    ".txt",
    ".md",
    ".json",
    ".jsonl",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".log",
}

EXTENSOES_TABELA = {
    ".csv",
    ".tsv",
}

EXTENSOES_CODIGO = {
    ".py",
    ".pyw",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".html",
    ".htm",
    ".css",
    ".scss",
    ".rs",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cs",
    ".java",
    ".kt",
    ".go",
    ".php",
    ".rb",
    ".sql",
    ".ps1",
    ".bat",
    ".cmd",
    ".sh",
    ".vue",
    ".svelte",
}

EXTENSOES_DOCUMENTO = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xlsm",
}

EXTENSOES = EXTENSOES_TEXTO | EXTENSOES_CODIGO | EXTENSOES_DOCUMENTO | EXTENSOES_TABELA


def resume_doc(texto: str, objetivo: str = "") -> str:
    from agent.config import KEEP_ALIVE, MODEL, NUM_CTX, create_client

    client = create_client()

    prompt = f"""Você é um componente auxiliar de análise.
    
    OBJETIVO:
    {objetivo or "Resumir o conteudo preservando informações importantes"}
    
    Analsie o conteúdo abaixo e devolva em blocos resumidos com as principais informações.
    Não invente informações ausentes no conteúdo.
    
    CONTEÚDO:
    {texto}
    """

    response = client.generate(
        model=MODEL,
        prompt=prompt,
        keep_alive=KEEP_ALIVE,
        options={"num_ctx": NUM_CTX},
    )

    return response.response


def _unidades_texto(caminho: Path):
    """Produz texto aos poucos; cada formato abre apenas o necessário.

    PDF extrai uma página por vez; DOCX percorre parágrafos; Excel usa o modo
    read_only. Textos e CSV são lidos progressivamente do arquivo.
    """
    extensao = caminho.suffix.lower()

    if extensao in EXTENSOES_TEXTO | EXTENSOES_CODIGO:
        with caminho.open("r", encoding="utf-8", errors="replace") as arquivo:
            while trecho := arquivo.read(4096):
                yield trecho

    elif extensao == ".pdf":
        import pymupdf

        with pymupdf.open(caminho) as documento:
            for numero, pagina in enumerate(documento, 1):
                yield f"\n[Página {numero}]\n" + pagina.get_text() + "\n"

    elif extensao == ".docx":
        # Lê o XML em fluxo; a biblioteca docx antiga presente no ambiente
        # também não precisa ser carregada para extrair texto.
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

        with ZipFile(caminho) as pacote, pacote.open("word/document.xml") as xml:
            eventos = ET.iterparse(xml, events=("start", "end"))
            _, raiz = next(eventos)

            for evento, elemento in eventos:
                if evento == "end" and elemento.tag == namespace + "p":
                    texto = "".join(
                        (fragmento.text or "")
                        if fragmento.tag == namespace + "t"
                        else "\t"
                        for fragmento in elemento.iter()
                        if fragmento.tag in {namespace + "t", namespace + "tab"}
                    )

                    if texto.strip():
                        yield texto + "\n"
                    raiz.clear()

    elif extensao in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        # read_only evita carregar todas as células e abas na memória.
        livro = load_workbook(caminho, read_only=True, data_only=True)
        try:
            for aba in livro.worksheets:
                yield f"\n[Aba: {aba.title}]\n"

                for numero, linha in enumerate(aba.iter_rows(values_only=True), 1):
                    valores = ["" if valor is None else str(valor) for valor in linha]
                    if any(valores):
                        yield f"Linha {numero}: " + " | ".join(valores) + "\n"

        finally:
            livro.close()

    elif extensao in {".csv", ".tsv"}:
        delimitador = "\t" if extensao == ".tsv" else ","

        with caminho.open(
            "r", encoding="utf-8-sig", errors="replace", newline=""
        ) as arquivo:
            leitor = csv.reader(arquivo, delimiter=delimitador)

            for numero, linha in enumerate(leitor, 1):
                yield f"Linha {numero}: " + " | ".join(linha) + "\n"


def _validar(caminho: str) -> Path:
    arquivo = Path(caminho).expanduser()

    if not arquivo.is_file():
        raise FileNotFoundError(f"Arquivo não encontrado: {arquivo}")

    if arquivo.suffix.lower() not in EXTENSOES:
        raise ValueError(f"Formato não suportado: {arquivo.suffix.lower()}")

    return arquivo


def ler_documento(caminho: str, inicio: int = 0, tamanho: int = TAMANHO_PADRAO) -> dict:
    """Lê até `tamanho` caracteres e informa onde continuar.

    inicio é um deslocamento em caracteres do texto extraído, não um número de
    página. O mesmo caminho e o proximo_inicio recuperam o trecho seguinte.
    """
    if inicio < 0 or not 1 <= tamanho <= TAMANHO_MAXIMO:
        return {"erro": f"inicio deve ser >= 0 e tamanho entre 1 e {TAMANHO_MAXIMO}."}
    try:
        arquivo = _validar(caminho)
        partes = []
        posicao = 0
        restante = tamanho + 1  # Um caractere extra distingue fim de continuação.

        for unidade in _unidades_texto(arquivo):
            fim = posicao + len(unidade)

            if fim <= inicio:
                posicao = fim
                continue

            trecho = unidade[max(0, inicio - posicao) :]
            partes.append(trecho[:restante])
            restante -= min(len(trecho), restante)
            posicao = fim

            if restante == 0:
                break

        texto = "".join(partes)
        conteudo = texto[:tamanho]
        tem_mais = len(texto) > tamanho

        if len(conteudo) >= LIMITE_RESUMO:
            conteudo_saida = resume_doc(conteudo)
        else:
            conteudo_saida = conteudo

        return {
            "caminho": str(arquivo.absolute()),
            "inicio": inicio,
            "conteudo": conteudo_saida,
            "proximo_inicio": inicio + len(conteudo) if tem_mais else None,
            "tem_mais": tem_mais,
        }

    except Exception as erro:
        return {"erro": f"Erro ao ler {caminho}: {type(erro).__name__}: {erro}"}


def normalizar_caminho(caminho: str) -> str:
    import re

    caminho = caminho.strip().strip('"').strip("'")

    # "Z" → "Z:/"
    if re.fullmatch(r"[A-Za-z]", caminho):
        return f"{caminho.upper()}:/"

    # "Z:" → "Z:/"
    if re.fullmatch(r"[A-Za-z]:", caminho):
        return f"{caminho.upper()}/"

    return caminho


def buscar_documento(caminho: str, termo: str, limite: int = 8) -> dict:
    """Busca trechos relevantes sem enviar o documento todo para o modelo.

    Faz busca textual sem diferenciar maiúsculas/minúsculas. Em arquivos imensos
    a varredura tem limite explícito e `busca_completa` indica o resultado parcial.
    """
    if not termo.strip() or not 1 <= limite <= 20:
        return {"erro": "Informe termo e limite entre 1 e 20."}
    try:
        caminho = normalizar_caminho(caminho)
        arquivo = _validar(caminho)
        termo_normalizado = termo.casefold()
        trechos = []
        posicao = 0
        cauda = ""
        ultimo_fim = 0

        for unidade in _unidades_texto(arquivo):
            texto = cauda + unidade
            base = posicao - len(cauda)
            inicio_busca = 0

            while len(trechos) < limite:
                indice = texto.casefold().find(termo_normalizado, inicio_busca)

                if indice < 0:
                    break
                # A cauda repete o final do bloco anterior para achar expressões
                # que cruzem a borda; não retornamos duas vezes a mesma ocorrência.

                if base + indice >= ultimo_fim:
                    trechos.append(
                        {
                            "inicio": base + indice,
                            "trecho": texto[
                                max(0, indice - 100) : indice + len(termo) + 140
                            ].strip(),
                        }
                    )

                    ultimo_fim = base + indice + len(termo)

                inicio_busca = indice + max(1, len(termo))

            posicao += len(unidade)

            if len(trechos) >= limite or posicao >= LIMITE_BUSCA_CARACTERES:
                return {
                    "caminho": str(arquivo.absolute()),
                    "termo": termo,
                    "trechos": trechos,
                    "busca_completa": False,
                    "aviso": "Limite de resultados ou caracteres atingido; pode haver mais ocorrências.",
                }

            cauda = texto[-max(0, len(termo) - 1) :] if len(termo) > 1 else ""

        return {
            "caminho": str(arquivo.absolute()),
            "termo": termo,
            "trechos": trechos,
            "busca_completa": True,
        }

    except Exception as erro:
        return {"erro": f"Erro ao buscar em {caminho}: {type(erro).__name__}: {erro}"}
